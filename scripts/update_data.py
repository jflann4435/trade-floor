#!/usr/bin/env python3
"""
Nightly data update for the Trade Floor class stock site.

What it does
  1. Reads the stock list from stocks/Trade_Stock_List.xlsx (rows with Status = Include).
  2. Pulls prices from Alpaca's market-data API:
       - 5 years of daily bars (charts, 52-week range, average volume, beta)
       - minute bars for the last 5 trading days (the replay and the 1D/5D charts)
       - bid/ask quotes sampled through the most recent trading day
       - cash dividends and each stock's exchange
  3. Pulls company numbers from SEC EDGAR filings: shares outstanding, EPS (TTM),
     and stockholders' equity (for book value per share).
  4. Writes data/summary.json (what every page needs) and data/history/<TICKER>.json
     (chart history, loaded when a stock is opened).

Usage
  python scripts/update_data.py           real data; needs the environment variables below
  python scripts/update_data.py --demo    made-up sample data, no keys or internet needed

Environment variables (set as GitHub repository secrets)
  ALPACA_KEY_ID, ALPACA_SECRET_KEY   from a free Alpaca paper-trading account
  SEC_CONTACT                        your name and email; the SEC requires it on every request
  ALPACA_FEED                        optional: "sip" (all U.S. exchanges, the default) or "iex"
"""
from __future__ import annotations

import argparse
import json
import math
import os
import random
import re
import shutil
import sys
import time
from collections import deque
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

ROOT = Path(__file__).resolve().parent.parent
STOCK_LIST = ROOT / "stocks" / "Trade_Stock_List.xlsx"
OUT_DIR = ROOT / "data"
ET = ZoneInfo("America/New_York")

# ---- settings you might change ----------------------------------------------
BENCHMARK = "SPY"            # shown on every page; also used for beta
YEARS_OF_HISTORY = 5
QUOTE_EVERY_MIN = 5          # sample bid/ask every N minutes of the replay day
QUOTE_SIZE_MULTIPLIER = 1    # if bid/ask sizes look 100x smaller than Yahoo's, set to 100
MIN_SUCCESS_SHARE = 0.9      # refuse to publish if fewer than 90% of stocks got prices
ETF_NAMES = {"SPY": "S&P 500", "XLP": "Consumer Staples", "XLI": "Industrials", "XLK": "Technology",
             "XHB": "Homebuilders", "XLY": "Consumer Discretionary", "XLB": "Materials",
             "XLE": "Energy", "XLF": "Financials", "XLU": "Utilities", "XLV": "Health Care",
             "ITA": "Aerospace & Defense", "ITB": "Home Construction"}
# -----------------------------------------------------------------------------

NOTES: list[str] = []


def note(msg: str) -> None:
    NOTES.append(msg)
    print("NOTE:", msg, flush=True)


def slugify(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


def r4(x):
    return None if x is None else round(float(x), 4)


# ============================================================================
# Stock list
# ============================================================================
def read_stock_list():
    from openpyxl import load_workbook

    wb = load_workbook(STOCK_LIST, data_only=True)

    # Categories tab -> trade buttons
    cs = wb["Categories"]
    hdr = [str(c.value).strip() if c.value is not None else "" for c in cs[1]]
    ci = {h: i for i, h in enumerate(hdr)}
    cats = []
    for row in cs.iter_rows(min_row=2, values_only=True):
        label = row[ci["Button Label"]] if "Button Label" in ci else None
        if not label or str(label).strip().lower() == "total":
            continue
        order = row[ci.get("Button Order", 0)]
        if not isinstance(order, (int, float)):
            continue
        cats.append({
            "order": order,
            "slug": slugify(str(label)),
            "name": str(label).strip(),
            "desc": str(row[ci["Button Description"]] or "").strip() if "Button Description" in ci else "",
            "etf": str(row[ci["Sector Benchmark"]] or BENCHMARK).strip().upper() if "Sector Benchmark" in ci else BENCHMARK,
        })
    cats.sort(key=lambda c: c["order"])
    for c in cats:
        del c["order"]
    by_name = {c["name"].lower(): c["slug"] for c in cats}

    # Stocks tab
    ws = wb["Stocks"]
    hdr = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    col = {h: i for i, h in enumerate(hdr)}
    for need in ("Ticker", "Company", "Category", "Status"):
        if need not in col:
            sys.exit(f"Stocks tab is missing the '{need}' column.")

    def g(row, name):
        i = col.get(name)
        return row[i] if i is not None and i < len(row) else None

    stocks = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        tk = str(g(row, "Ticker") or "").strip().upper()
        if not tk or str(g(row, "Status") or "").strip() != "Include":
            continue
        cat = by_name.get(str(g(row, "Category") or "").strip().lower())
        if not cat:
            note(f"{tk}: category '{g(row, 'Category')}' isn't on the Categories tab, so it was skipped.")
            continue
        also = []
        for part in str(g(row, "Also In") or "").split(";"):
            s = by_name.get(part.strip().lower())
            if s and s != cat:
                also.append(s)
        target = g(row, "1y Target Est")
        try:
            target = float(target) if target not in (None, "") else None
        except (TypeError, ValueError):
            note(f"{tk}: '1y Target Est' isn't a number ({target!r}); ignored.")
            target = None
        asof = g(row, "Target As Of")
        if isinstance(asof, datetime):
            asof = asof.date().isoformat()
        elif isinstance(asof, date):
            asof = asof.isoformat()
        elif asof:
            asof = str(asof)
        stocks[tk] = {
            "company": str(g(row, "Company") or tk).strip(),
            "category": cat,
            "also_in": also,
            "role": str(g(row, "Role") or "").strip(),
            "what": str(g(row, "What They Do") or "").strip(),
            "hook": str(g(row, "Student Hook") or "").strip(),
            "target": target,
            "target_asof": asof or None,
        }
    if not stocks:
        sys.exit("No rows with Status = Include were found on the Stocks tab.")
    etfs = sorted({BENCHMARK} | {c["etf"] for c in cats})
    return cats, stocks, etfs


# ============================================================================
# Alpaca
# ============================================================================
class ApiError(Exception):
    def __init__(self, status, text):
        super().__init__(f"HTTP {status}: {text}")
        self.status = status
        self.text = text


class Alpaca:
    DATA = "https://data.alpaca.markets"
    TRADE = "https://paper-api.alpaca.markets"
    MAX_PER_MIN = 180  # free plan allows 200 historical calls per minute

    def __init__(self, key, secret, feed):
        import requests

        self.s = requests.Session()
        self.s.headers.update({"APCA-API-KEY-ID": key, "APCA-API-SECRET-KEY": secret})
        self.feed = feed
        self.recent = deque()
        self.calls = 0

    def _throttle(self):
        now = time.monotonic()
        while self.recent and now - self.recent[0] > 60:
            self.recent.popleft()
        if len(self.recent) >= self.MAX_PER_MIN:
            time.sleep(60 - (now - self.recent[0]) + 0.2)
        self.recent.append(time.monotonic())

    def get(self, url, params=None):
        for attempt in range(6):
            self._throttle()
            self.calls += 1
            try:
                r = self.s.get(url, params=params, timeout=60)
            except Exception as e:  # network hiccup
                time.sleep(3 * (attempt + 1))
                last = e
                continue
            if r.status_code == 429 or r.status_code >= 500:
                time.sleep(5 * (attempt + 1))
                last = ApiError(r.status_code, r.text[:200])
                continue
            if r.status_code != 200:
                raise ApiError(r.status_code, r.text[:300])
            return r.json()
        raise last

    def _data_get(self, url, params):
        """GET with the feed parameter; if the account can't use SIP, fall back to IEX once."""
        p = dict(params, feed=self.feed)
        try:
            return self.get(url, p)
        except ApiError as e:
            if self.feed == "sip" and e.status in (403, 422) and "subscription" in e.text.lower():
                note("This Alpaca account can't read SIP (all-exchange) data, so the IEX feed is used. "
                     "Prices will be close to Yahoo's, but volume will look much smaller.")
                self.feed = "iex"
                return self.get(url, dict(params, feed="iex"))
            raise

    def bars(self, symbols, timeframe, start, end, adjustment):
        out = {s: [] for s in symbols}
        for i in range(0, len(symbols), 25):
            chunk = symbols[i:i + 25]
            params = {"symbols": ",".join(chunk), "timeframe": timeframe, "start": iso_z(start),
                      "end": iso_z(end), "limit": 10000, "adjustment": adjustment, "sort": "asc"}
            while True:
                j = self._data_get(f"{self.DATA}/v2/stocks/bars", params)
                for sym, rows in (j.get("bars") or {}).items():
                    out.setdefault(sym, []).extend(rows)
                tok = j.get("next_page_token")
                if not tok:
                    break
                params["page_token"] = tok
        return out

    def first_good_quote(self, sym, start, end):
        j = self._data_get(f"{self.DATA}/v2/stocks/{sym}/quotes",
                           {"start": iso_z(start), "end": iso_z(end), "limit": 10})
        for q in j.get("quotes") or []:
            bp, ap = q.get("bp") or 0, q.get("ap") or 0
            if bp > 0 and ap > 0 and ap >= bp:
                return q
        return None

    def cash_dividends(self, symbols, start, end):
        out = {s: [] for s in symbols}
        params = {"symbols": ",".join(symbols), "types": "cash_dividend", "start": start.isoformat(),
                  "end": end.isoformat(), "limit": 1000}
        while True:
            j = self.get(f"{self.DATA}/v1/corporate-actions", params)
            ca = j.get("corporate_actions") or {}
            for d in ca.get("cash_dividends") or []:
                if d.get("symbol") in out:
                    out[d["symbol"]].append(d)
            tok = j.get("next_page_token")
            if not tok:
                break
            params["page_token"] = tok
        return out

    def exchange(self, sym):
        j = self.get(f"{self.TRADE}/v2/assets/{sym}")
        ex = (j.get("exchange") or "").upper()
        return {"NASDAQ": "Nasdaq", "NYSE": "NYSE", "ARCA": "NYSE Arca", "AMEX": "NYSE American",
                "BATS": "Cboe BZX"}.get(ex, ex or "?")


def iso_z(dt: datetime) -> str:
    return dt.astimezone(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def et_dt(t: str) -> datetime:
    return datetime.fromisoformat(t.replace("Z", "+00:00")).astimezone(ET)


def session_open(d: date) -> datetime:
    return datetime(d.year, d.month, d.day, 9, 30, tzinfo=ET)


# ============================================================================
# SEC EDGAR
# ============================================================================
class Sec:
    def __init__(self, contact):
        import requests

        self.s = requests.Session()
        self.s.headers.update({"User-Agent": f"TradeFloor classroom stock site {contact}",
                               "Accept-Encoding": "gzip, deflate"})

    def get(self, url):
        for attempt in range(4):
            time.sleep(0.15)  # SEC asks for no more than 10 requests per second
            r = self.s.get(url, timeout=60)
            if r.status_code == 200:
                return r.json()
            if r.status_code in (429, 500, 502, 503):
                time.sleep(3 * (attempt + 1))
                continue
            raise ApiError(r.status_code, r.text[:200])
        raise ApiError(r.status_code, "SEC kept refusing the request")

    def cik_map(self):
        j = self.get("https://www.sec.gov/files/company_tickers.json")
        return {v["ticker"].upper().replace("-", "."): int(v["cik_str"]) for v in j.values()}

    def facts(self, cik):
        return self.get(f"https://data.sec.gov/api/xbrl/companyfacts/CIK{cik:010d}.json")


def fact_entries(facts, concepts, unit, taxonomy="us-gaap"):
    for c in concepts:
        u = facts.get("facts", {}).get(taxonomy, {}).get(c, {}).get("units", {}).get(unit)
        if u:
            return u
    return []


def _periodic(e):
    return str(e.get("form", "")).startswith(("10-K", "10-Q"))


def latest_instant(entries):
    """Latest point-in-time value (e.g. equity on the balance sheet). Returns (end_date, [values])."""
    ok = [e for e in entries if _periodic(e)] or entries
    if not ok:
        return None, []
    end = max(e["end"] for e in ok)
    at = [e for e in ok if e["end"] == end]
    filed = max(e.get("filed", "") for e in at)
    vals, seen = [], set()
    for e in at:
        if e.get("filed", "") == filed and (e.get("accn"), e["val"]) not in seen:
            seen.add((e.get("accn"), e["val"]))
            vals.append(e["val"])
    return end, vals


def ttm(entries):
    """Trailing-twelve-month total from 10-K/10-Q facts: last fiscal year + this year-to-date - last year's same YTD."""
    rows = {}
    for e in entries:
        if "start" not in e or not _periodic(e):
            continue
        k = (e["start"], e["end"])
        if k not in rows or e.get("filed", "") > rows[k].get("filed", ""):
            rows[k] = e
    if not rows:
        return None, None
    items = [(date.fromisoformat(s), date.fromisoformat(en), e["val"]) for (s, en), e in rows.items()]
    dur = lambda s, en: (en - s).days
    latest = max(en for _, en, _ in items)
    if (date.today() - latest).days > 460:
        return None, None  # too old to trust
    annual = [x for x in items if 350 <= dur(x[0], x[1]) <= 380]
    for s, en, v in annual:
        if en == latest:
            return v, latest.isoformat()
    ytd = [x for x in items if x[1] == latest and dur(x[0], x[1]) < 350]
    if ytd:
        s0, en0, v0 = max(ytd, key=lambda x: dur(x[0], x[1]))
        prev_fy = [a for a in annual if abs((a[1] - (s0 - timedelta(days=1))).days) <= 10]
        prior = [p for p in items if abs((p[1] - (en0 - timedelta(days=365))).days) <= 10
                 and abs(dur(p[0], p[1]) - dur(s0, en0)) <= 10]
        if prev_fy and prior:
            return prev_fy[0][2] + v0 - prior[0][2], latest.isoformat()
    quarters = sorted([x for x in items if 80 <= dur(x[0], x[1]) <= 100], key=lambda x: x[1])
    if len(quarters) >= 4 and (quarters[-1][1] - quarters[-4][1]).days <= 290:
        return sum(x[2] for x in quarters[-4:]), latest.isoformat()
    if annual:
        a = max(annual, key=lambda x: x[1])
        return a[2], a[1].isoformat()
    return None, None


def sec_fundamentals(sec, cik_map, tk):
    cik = cik_map.get(tk)
    if not cik:
        note(f"{tk}: not found in the SEC's ticker list; market cap, EPS and book value will be blank.")
        return {}
    f = sec.facts(cik)
    out = {}
    end, vals = latest_instant(fact_entries(f, ["EntityCommonStockSharesOutstanding"], "shares", "dei"))
    if not vals:
        end, vals = latest_instant(fact_entries(f, ["CommonStockSharesOutstanding"], "shares"))
    if vals:
        out["shares"] = float(sum(vals))
        out["shares_asof"] = end
    eps, eps_end = ttm(fact_entries(f, ["EarningsPerShareDiluted", "EarningsPerShareBasic"], "USD/shares"))
    if eps is not None:
        out["eps"] = eps
        out["eps_asof"] = eps_end
    end, vals = latest_instant(fact_entries(
        f, ["StockholdersEquity", "StockholdersEquityIncludingPortionAttributableToNoncontrollingInterest"], "USD"))
    if vals and out.get("shares"):
        out["bvps"] = vals[0] / out["shares"]
        out["book_asof"] = end
    dps, _ = ttm(fact_entries(f, ["CommonStockDividendsPerShareDeclared", "CommonStockDividendsPerShareCashPaid"],
                              "USD/shares"))
    if dps is not None:
        out["div_sec"] = dps
    return out


# ============================================================================
# Calculations
# ============================================================================
def beta_5y(adj_daily, bench_daily):
    """Beta from up to 60 monthly returns vs. the benchmark (the same method Yahoo labels '5Y Monthly')."""
    def month_closes(rows):
        m = {}
        for d, c in rows:
            m[(d.year, d.month)] = c
        return m
    a, b = month_closes(adj_daily), month_closes(bench_daily)
    months = sorted(set(a) & set(b))[-61:]
    if len(months) < 25:
        return None
    ra = [a[months[i]] / a[months[i - 1]] - 1 for i in range(1, len(months))]
    rb = [b[months[i]] / b[months[i - 1]] - 1 for i in range(1, len(months))]
    ma, mb = sum(ra) / len(ra), sum(rb) / len(rb)
    cov = sum((x - ma) * (y - mb) for x, y in zip(ra, rb)) / (len(ra) - 1)
    var = sum((y - mb) ** 2 for y in rb) / (len(rb) - 1)
    return cov / var if var else None


def forward_dividend(divs, as_of: date):
    """Latest regular dividend x how many were paid in the past year (Yahoo's 'forward dividend')."""
    regular = [d for d in divs if not d.get("special") and d.get("rate")]
    regular.sort(key=lambda d: d.get("ex_date") or "")
    if not regular:
        return 0.0
    last = regular[-1]
    try:
        if (as_of - date.fromisoformat(last["ex_date"])).days > 400:
            return 0.0  # stopped paying
    except Exception:
        pass
    yr = [d for d in regular if d.get("ex_date") and (as_of - date.fromisoformat(d["ex_date"])).days <= 365]
    n = min(12, max(1, len(yr)))
    return float(last["rate"]) * n


def build_session(rows, day: date):
    """Turn one day's minute bars into a full session (one row per minute, gaps carried forward)."""
    open_dt = session_open(day)
    by_min = {}
    for b in rows:
        t = et_dt(b["t"])
        if t.date() != day:
            continue
        m = int((t - open_dt).total_seconds() // 60)
        if 0 <= m < 390:
            by_min[m] = b
    return by_min


def session_length(spy_by_min):
    if not spy_by_min:
        return 390
    return 210 if max(spy_by_min) < 215 else 390  # early-close days end at 1:00 PM


def fill_session(by_min, length):
    out, last = [], None
    first = by_min[min(by_min)] if by_min else None
    for m in range(length):
        b = by_min.get(m)
        if b:
            row = [b["o"], b["h"], b["l"], b["c"], int(b["v"])]
            last = b["c"]
        else:
            p = last if last is not None else (first["o"] if first else None)
            if p is None:
                return None
            row = [p, p, p, p, 0]
        out.append([r4(x) for x in row[:4]] + [row[4]])
    return out


def five_minute(session):
    out = []
    for i in range(0, len(session), 5):
        ch = session[i:i + 5]
        out.append([ch[0][0], max(x[1] for x in ch), min(x[2] for x in ch), ch[-1][3], sum(x[4] for x in ch)])
    return out


# ============================================================================
# Real run
# ============================================================================
def run_real(cats, stocks, etfs):
    key, secret = os.environ.get("ALPACA_KEY_ID"), os.environ.get("ALPACA_SECRET_KEY")
    contact = os.environ.get("SEC_CONTACT", "").strip()
    if not key or not secret:
        sys.exit("ALPACA_KEY_ID and ALPACA_SECRET_KEY must be set (repository secrets).")
    if not contact:
        sys.exit("SEC_CONTACT must be set to your name and email (the SEC requires it).")
    al = Alpaca(key, secret, os.environ.get("ALPACA_FEED", "sip").lower())
    tickers = list(stocks)
    symbols = sorted(set(tickers) | set(etfs))

    now_et = datetime.now(ET)
    # Never ask for the last 15 minutes (free plan rule). During market hours, stop at yesterday.
    end = now_et - timedelta(minutes=16)
    if now_et.weekday() < 5 and now_et.time() < datetime.strptime("16:20", "%H:%M").time():
        end = datetime(now_et.year, now_et.month, now_et.day, tzinfo=ET) - timedelta(seconds=1)
    start = end - timedelta(days=365 * YEARS_OF_HISTORY + 10)

    print("Daily bars…", flush=True)
    daily = al.bars(symbols, "1Day", start, end, "split")
    daily_adj = al.bars(symbols, "1Day", start, end, "all")
    to_rows = lambda rows: [(et_dt(b["t"]).date(), b) for b in rows]
    spy_days = [d for d, _ in to_rows(daily.get(BENCHMARK, []))]
    if len(spy_days) < 300:
        sys.exit(f"Only {len(spy_days)} days of {BENCHMARK} data came back; not publishing.")
    replay_day = spy_days[-1]
    last5 = spy_days[-5:]
    print(f"Replay day: {replay_day}", flush=True)

    print("Minute bars…", flush=True)
    minute = al.bars(symbols, "1Min", session_open(last5[0]),
                     datetime(replay_day.year, replay_day.month, replay_day.day, 16, 0, tzinfo=ET), "split")
    spy_sessions = {d: build_session(minute.get(BENCHMARK, []), d) for d in last5}
    lengths = {d: session_length(spy_sessions[d]) for d in last5}
    L = lengths[replay_day]

    print("Exchanges and dividends…", flush=True)
    exchanges = {}
    for s in symbols:
        try:
            exchanges[s] = al.exchange(s)
        except Exception as e:
            note(f"{s}: couldn't look up the exchange ({e}).")
            exchanges[s] = "?"
    try:
        divs = al.cash_dividends(tickers + etfs, replay_day - timedelta(days=400), replay_day)
        divs_ok = True
    except Exception as e:
        note(f"Alpaca dividend lookup failed ({e}); using SEC filings for dividends instead.")
        divs, divs_ok = {}, False

    print("SEC filings…", flush=True)
    sec = Sec(contact)
    try:
        cik_map = sec.cik_map()
    except Exception as e:
        note(f"SEC ticker list failed ({e}); company numbers will be blank tonight.")
        cik_map = {}
    fund = {}
    for tk in tickers:
        try:
            fund[tk] = sec_fundamentals(sec, cik_map, tk) if cik_map else {}
        except Exception as e:
            note(f"{tk}: SEC data failed ({e}).")
            fund[tk] = {}

    print("Bid/ask quotes…", flush=True)
    samples = [-5] + list(range(0, L, QUOTE_EVERY_MIN)) + [L - 1, L + 5]
    open_dt = session_open(replay_day)

    bench_adj = [(d, b["c"]) for d, b in to_rows(daily_adj.get(BENCHMARK, []))]
    summary_syms, histories, ok = {}, {}, 0
    for s in symbols:
        rows = [(d, b) for d, b in to_rows(daily.get(s, [])) if d <= replay_day]
        sessions = {d: build_session(minute.get(s, []), d) for d in last5}
        today = fill_session(sessions[replay_day], L) if sessions[replay_day] else None
        if not rows or rows[-1][0] != replay_day or today is None:
            note(f"{s}: no prices for {replay_day}; it will be missing from the site tonight.")
            continue
        ok += s in stocks
        before = rows[:-1]
        prev = before[-1][1]["c"]
        last_year = before[-251:]
        entry = {
            "prev": r4(prev),
            "wk_lo": r4(min(b["l"] for _, b in last_year)),
            "wk_hi": r4(max(b["h"] for _, b in last_year)),
            "avg_vol": int(sum(b["v"] for _, b in before[-63:]) / max(1, len(before[-63:]))),
            "today": today,
        }
        if s in stocks:
            quotes = []
            for m in samples:
                t0 = open_dt + timedelta(minutes=m)
                try:
                    q = al.first_good_quote(s, t0, t0 + timedelta(minutes=5))
                except Exception as e:
                    note(f"{s}: quote lookup failed at minute {m} ({e}).")
                    q = None
                if q:
                    quotes.append([m, r4(q["bp"]), int(q.get("bs", 0) * QUOTE_SIZE_MULTIPLIER),
                                   r4(q["ap"]), int(q.get("as", 0) * QUOTE_SIZE_MULTIPLIER)])
            entry["quotes"] = quotes
            f = fund.get(s, {})
            div = forward_dividend(divs.get(s, []), replay_day) if divs_ok else f.get("div_sec", 0.0)
            adj = [(d, b["c"]) for d, b in to_rows(daily_adj.get(s, [])) if d <= replay_day]
            entry.update({
                "shares": f.get("shares"), "eps": r4(f.get("eps")), "bvps": r4(f.get("bvps")),
                "eps_asof": f.get("eps_asof"), "book_asof": f.get("book_asof"),
                "div": r4(div or 0.0), "beta": r4(beta_5y(adj, bench_adj)),
                "target": stocks[s]["target"], "target_asof": stocks[s]["target_asof"],
            })
        summary_syms[s] = entry
        hist_days = [d for d in last5[:-1]]
        histories[s] = {
            "daily": [[d.isoformat(), r4(b["o"]), r4(b["h"]), r4(b["l"]), r4(b["c"]), int(b["v"])] for d, b in before],
            "fivemin": [{"date": d.isoformat(), "bars": five_minute(fill_session(sessions[d], lengths[d]) or [])}
                        for d in hist_days if sessions[d]],
        }

    if ok < MIN_SUCCESS_SHARE * len(stocks):
        sys.exit(f"Only {ok} of {len(stocks)} stocks got prices; keeping last night's data.")
    if BENCHMARK not in summary_syms:
        sys.exit(f"{BENCHMARK} is missing; keeping last night's data.")
    print(f"Alpaca calls: {al.calls}", flush=True)
    return {"replay_day": replay_day.isoformat(), "session_len": L, "feed": al.feed,
            "exchanges": exchanges, "symbols": summary_syms}, histories


# ============================================================================
# Demo run (offline, made-up numbers)
# ============================================================================
DEMO_LEVELS = {"SPY": 680, "XLP": 80, "XLI": 150, "XLK": 280, "XHB": 105, "XLY": 230}


def run_demo(cats, stocks, etfs):
    today = datetime.now(ET).date()
    d = today - timedelta(days=1)
    while d.weekday() >= 5:
        d -= timedelta(days=1)
    days = []
    x = d
    while len(days) < 1265:
        if x.weekday() < 5:
            days.append(x)
        x -= timedelta(days=1)
    days.reverse()
    replay_day, last5 = days[-1], days[-5:]
    syms, hists = {}, {}
    for s in sorted(set(stocks) | set(etfs)):
        rnd = random.Random(s)
        base = DEMO_LEVELS.get(s, rnd.uniform(20, 600))
        beta = rnd.uniform(0.5, 1.8)
        sig = 0.006 + 0.008 * beta
        vol = rnd.uniform(0.3, 20) * 1e6
        # backwards daily walk from today's level
        closes = [base]
        for _ in range(len(days) - 1):
            closes.append(closes[-1] / (1 + 0.0003 + rnd.gauss(0, sig * 0.75)))
        closes.reverse()
        daily = []
        for i, day in enumerate(days):
            c = closes[i]
            o = c / (1 + rnd.gauss(0, sig * 0.6))
            daily.append([day.isoformat(), r4(o), r4(max(o, c) * (1 + abs(rnd.gauss(0, sig * .4)))),
                          r4(min(o, c) * (1 - abs(rnd.gauss(0, sig * .4)))), r4(c), int(vol * rnd.uniform(.6, 1.4))])

        def session(open_px):
            rows, p = [], open_px
            sm = sig * 0.9 / math.sqrt(390)
            for m in range(390):
                c = p * (1 + rnd.gauss(0, sm))
                w = 1 + 2.4 * math.exp(-m / 22) + 1.8 * math.exp(-(389 - m) / 18)
                rows.append([r4(p), r4(max(p, c) * (1 + abs(rnd.gauss(0, sm * .5)))),
                             r4(min(p, c) * (1 - abs(rnd.gauss(0, sm * .5)))), r4(c), int(vol / 390 * w * rnd.uniform(.5, 1.2))])
                p = c
            return rows
        sessions = {day: session(daily[days.index(day) - 1][4] * (1 + rnd.gauss(0, sig * .3))) for day in last5}
        before = daily[:-1]
        prev = before[-1][4]
        e = {"prev": prev, "wk_lo": min(r[3] for r in before[-251:]), "wk_hi": max(r[2] for r in before[-251:]),
             "avg_vol": int(sum(r[5] for r in before[-63:]) / 63), "today": sessions[replay_day]}
        if s in stocks:
            spread = max(0.01, base * rnd.choice([0.0001, 0.0004, 0.001]))
            qs = []
            for m in [-5] + list(range(0, 390, QUOTE_EVERY_MIN)) + [389, 395]:
                px = sessions[replay_day][min(max(m, 0), 389)][3]
                sp = spread * (6 if m < 0 or m > 389 else rnd.uniform(.7, 2))
                bid = round(px - sp / 2, 2)
                qs.append([m, bid, rnd.randint(1, 12) * 100, round(bid + max(.01, round(sp, 2)), 2), rnd.randint(1, 12) * 100])
            pe = rnd.choice([12, 18, 25, 35, 60, -80])
            e.update({"quotes": qs, "shares": rnd.uniform(.05, 3) * 1e9, "eps": r4(base / pe),
                      "bvps": r4(base / rnd.uniform(1, 15)), "div": r4(base * rnd.choice([0, 0, .01, .02, .035])),
                      "beta": r4(beta), "target": stocks[s]["target"] or r4(base * rnd.uniform(1, 1.18)),
                      "target_asof": stocks[s]["target_asof"] or today.isoformat(),
                      "eps_asof": None, "book_asof": None})
        syms[s] = e
        hists[s] = {"daily": before, "fivemin": [{"date": day.isoformat(), "bars": five_minute(sessions[day])}
                                                  for day in last5[:-1]]}
    exchanges = {s: ("Nasdaq" if random.Random(s + "x").random() < .3 else "NYSE") for s in syms}
    for e in etfs:
        exchanges[e] = "NYSE Arca"
    return {"replay_day": replay_day.isoformat(), "session_len": 390, "feed": "demo",
            "exchanges": exchanges, "symbols": syms}, hists


# ============================================================================
def write_output(cats, stocks, etfs, core, histories, demo):
    tmp = ROOT / "data_tmp"
    if tmp.exists():
        shutil.rmtree(tmp)
    (tmp / "history").mkdir(parents=True)
    ex = core.pop("exchanges")
    summary = {
        "version": 1,
        "demo": demo,
        "generated": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        **core,
        "categories": cats,
        "stocks": {tk: {**{k: v for k, v in m.items() if k not in ("target", "target_asof")},
                        "exchange": ex.get(tk, "?")} for tk, m in stocks.items() if tk in core["symbols"]},
        "etfs": {e: {"name": ETF_NAMES.get(e, e), "exchange": ex.get(e, "?")} for e in etfs if e in core["symbols"]},
        "benchmark": BENCHMARK,
        "notes": NOTES,
    }
    (tmp / "summary.json").write_text(json.dumps(summary, separators=(",", ":")))
    for s, h in histories.items():
        (tmp / "history" / f"{s}.json").write_text(json.dumps(h, separators=(",", ":")))
    if OUT_DIR.exists():
        shutil.rmtree(OUT_DIR)
    tmp.rename(OUT_DIR)
    print(f"Wrote {OUT_DIR / 'summary.json'} and {len(histories)} history files.")

    gh = os.environ.get("GITHUB_STEP_SUMMARY")
    if gh:
        with open(gh, "a") as f:
            f.write(f"### Trade Floor data for {core['replay_day']}\n\n")
            f.write(f"- Stocks on the site: {len(summary['stocks'])} (feed: {core['feed']})\n")
            f.write("".join(f"- {n}\n" for n in NOTES) or "- No problems found.\n")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--demo", action="store_true", help="write made-up sample data (no keys needed)")
    args = ap.parse_args()
    cats, stocks, etfs = read_stock_list()
    print(f"{len(stocks)} stocks, {len(cats)} trades, benchmarks {', '.join(etfs)}", flush=True)
    core, hist = run_demo(cats, stocks, etfs) if args.demo else run_real(cats, stocks, etfs)
    write_output(cats, stocks, etfs, core, hist, args.demo)


if __name__ == "__main__":
    main()
